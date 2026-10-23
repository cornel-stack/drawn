"""Import the ``Backlog`` sheet of each channel's spreadsheet into backlog.yaml.

The three workbooks use different overall column sets but share the seven
columns we need, under identical header names. We map by header name (not
position) so the importer is robust to the differing layouts.

* ``target_minutes`` — the midpoint of a range like ``"6-8 min"`` (→ 7), or the
  single value of ``"12 min"`` (→ 12).
* ``prereqs`` — comma-split; ``"None"``/blank → ``[]``; range tokens such as
  ``"EM-01…EM-08"`` are kept verbatim as one entry.
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from pathlib import Path

import openpyxl
import yaml

CHANNELS = ("chalk-and-vector", "pen-and-pendulum", "marginal-ink")

# backlog.yaml field -> spreadsheet column header
COLUMNS = {
    "code": "Code",
    "title": "Working title",
    "course": "Course",
    "source": "Source section (textbook)",
    "format": "Format",
    "target_minutes": "Target length",
    "prereqs": "Prerequisites",
}

FIELD_ORDER = ("code", "title", "course", "source", "format", "target_minutes", "prereqs")


def parse_target_minutes(value: object) -> int | None:
    if value is None:
        return None
    nums = [int(n) for n in re.findall(r"\d+", str(value))]
    if not nums:
        return None
    return round(sum(nums) / len(nums))


def parse_prereqs(value: object) -> list[str]:
    if value is None:
        return []
    text = str(value).strip()
    if not text or text.lower() == "none":
        return []
    return [part.strip() for part in text.split(",") if part.strip()]


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def find_backlog_xlsx(research_dir: Path) -> Path | None:
    hits = sorted(research_dir.glob("*backlog*.xlsx"))
    return hits[0] if hits else None


def import_rows(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Backlog" not in wb.sheetnames:
        raise ValueError(f"{xlsx_path}: no 'Backlog' sheet (found {wb.sheetnames})")
    ws = wb["Backlog"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {name: i for i, name in enumerate(header)}
    missing = [col for col in COLUMNS.values() if col not in idx]
    if missing:
        raise ValueError(f"{xlsx_path}: missing columns {missing}")

    out: list[dict] = []
    for row in rows[1:]:
        code = _clean(row[idx[COLUMNS["code"]]])
        if not code:
            continue
        out.append(
            {
                "code": code,
                "title": _clean(row[idx[COLUMNS["title"]]]),
                "course": _clean(row[idx[COLUMNS["course"]]]),
                "source": _clean(row[idx[COLUMNS["source"]]]),
                "format": _clean(row[idx[COLUMNS["format"]]]),
                "target_minutes": parse_target_minutes(row[idx[COLUMNS["target_minutes"]]]),
                "prereqs": parse_prereqs(row[idx[COLUMNS["prereqs"]]]),
            }
        )
    return out


def write_backlog_yaml(rows: Iterable[dict], out_path: Path) -> None:
    ordered = [{field: row[field] for field in FIELD_ORDER} for row in rows]
    header = (
        "# Generated from research/*_backlog.xlsx by `just fetch`.\n"
        "# Do not hand-edit; re-run the importer instead.\n"
    )
    out_path.write_text(
        header
        + yaml.safe_dump(ordered, sort_keys=False, allow_unicode=True, width=100)
    )


def import_channel(channel_dir: Path) -> int:
    research = channel_dir / "research"
    xlsx = find_backlog_xlsx(research)
    if xlsx is None:
        raise FileNotFoundError(f"no *backlog*.xlsx under {research}")
    rows = import_rows(xlsx)
    write_backlog_yaml(rows, channel_dir / "backlog.yaml")
    return len(rows)


def import_all(repo_root: Path) -> dict[str, int]:
    counts: dict[str, int] = {}
    for name in CHANNELS:
        channel_dir = repo_root / "channels" / name
        if channel_dir.is_dir():
            counts[name] = import_channel(channel_dir)
    return counts
